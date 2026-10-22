import docx
import re
import os

def parse_inline_markdown(para, text):
    # Separa por **negrito** e `codigo`
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if not token: 
            continue
        if token.startswith('**') and token.endswith('**'):
            run = para.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = para.add_run(token[1:-1])
        else:
            para.add_run(token)

def convert_md_to_docx(md_path, docx_path):
    doc = docx.Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
        
    paragraphs = content.split('\n')
    current_para = []
    
    def flush_para():
        nonlocal current_para
        if current_para:
            text = ' '.join(current_para)
            p = doc.add_paragraph()
            parse_inline_markdown(p, text)
            current_para = []

    for line in paragraphs:
        line = line.strip()
        
        # Pular linhas vazias e divisórias
        if not line or line == '---':
            flush_para()
            continue
            
        if line.startswith('# '):
            flush_para()
            doc.add_heading(line.replace('# ', '').strip(), level=1)
        elif line.startswith('## '):
            flush_para()
            doc.add_heading(line.replace('## ', '').strip(), level=2)
        elif line.startswith('### '):
            flush_para()
            doc.add_heading(line.replace('### ', '').strip(), level=3)
        elif line.startswith('> '):
            flush_para()
            text = line.replace('> ', '').strip()
            p = doc.add_paragraph()
            # Deixar itálico
            tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
            for token in tokens:
                if not token: continue
                if token.startswith('**') and token.endswith('**'):
                    run = p.add_run(token[2:-2])
                    run.bold = True
                    run.italic = True
                else:
                    run = p.add_run(token)
                    run.italic = True
        elif line.startswith('- ') or line.startswith('* '):
            flush_para()
            text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            parse_inline_markdown(p, text)
        elif re.match(r'^\d+\.\s*', line):
            flush_para()
            # Remove a numeração hardcoded (ex: "1. ") porque o Word já põe o número
            text = re.sub(r'^\d+\.\s*', '', line)
            p = doc.add_paragraph(style='List Number')
            parse_inline_markdown(p, text)
        else:
            current_para.append(line)
            
    flush_para()
                
    doc.save(docx_path)
    print("Success")

if __name__ == "__main__":
    md = r"C:\apps\IAProfEPT\PLATAFORMA\RELATORIO_DESENVOLVIMENTO.md"
    docx_file = r"C:\apps\IAProfEPT\PLATAFORMA\RELATORIO_DESENVOLVIMENTO_V2.docx"
    convert_md_to_docx(md, docx_file)
