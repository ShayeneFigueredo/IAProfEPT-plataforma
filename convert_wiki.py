import os
import docx

def main():
    wiki_dir = r"C:\apps\IAProfEPT\PLATAFORMA\docs\wiki"
    output_path = r"C:\apps\IAProfEPT\PLATAFORMA\Documentacao_Wiki_IAProfEPT.docx"
    
    # Get all .md files
    files = [f for f in os.listdir(wiki_dir) if f.endswith('.md')]
    # Sort files
    def sort_key(f):
        if f == "Home.md": return 0
        if "DIARIO" in f: return 1
        if "GUIA" in f: return 2
        if "MATRIZ" in f: return 3
        return 4 + files.index(f) if f in files else 99
        
    files.sort(key=lambda x: (sort_key(x), x))
    
    doc = docx.Document()
    doc.add_heading('Documentação do Projeto - IAProfEPT', 0)
    
    for f in files:
        filepath = os.path.join(wiki_dir, f)
        with open(filepath, 'r', encoding='utf-8') as file:
            md_text = file.read()
            doc.add_heading(f"Arquivo: {f}", level=1)
            
            # Simple text addition, splitting by double newlines for paragraphs
            paragraphs = md_text.split('\n\n')
            for p in paragraphs:
                if p.startswith('# '):
                    doc.add_heading(p.replace('# ', ''), level=1)
                elif p.startswith('## '):
                    doc.add_heading(p.replace('## ', ''), level=2)
                elif p.startswith('### '):
                    doc.add_heading(p.replace('### ', ''), level=3)
                else:
                    doc.add_paragraph(p)
                    
            doc.add_page_break()
    
    doc.save(output_path)
    print(f"Sucesso! Salvo em: {output_path}")

if __name__ == '__main__':
    main()

